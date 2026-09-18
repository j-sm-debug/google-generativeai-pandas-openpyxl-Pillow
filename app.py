import streamlit as st
import google.generativeai as genai
import PIL.Image
import pandas as pd
import io
import json
import os
import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# --- ページ設定 ---
st.set_page_config(page_title="売上日計表 自動照合システム", layout="wide")

# --- セッション状態の初期化 ---
if "accumulated_results" not in st.session_state:
    st.session_state["accumulated_results"] = pd.DataFrame()
if "accumulated_unreadable" not in st.session_state:
    st.session_state["accumulated_unreadable"] = pd.DataFrame()
if "file_uploader_key" not in st.session_state:
    st.session_state["file_uploader_key"] = 0

# --- 116店舗マスタの固定定義（コード内埋め込み） ---
STORE_MASTER_DATA = [
    {"店舗コード": "B75", "店舗名": "和光店"},
    {"店舗コード": "B79", "店舗名": "柏沼南店"},
    {"店舗コード": "B82", "店舗名": "船橋薬円台店"},
    {"店舗コード": "B83", "店舗名": "新取手店"},
    {"店舗コード": "B86", "店舗名": "石下店"},
    {"店舗コード": "B87", "店舗名": "野田店"},
    {"店舗コード": "B88", "店舗名": "葛飾白鳥店"},
    {"店舗コード": "B91", "店舗名": "鹿骨店"},
    {"店舗コード": "B95", "店舗名": "欠真間店"},
    {"店舗コード": "B96", "店舗名": "南流山店"},
    {"店舗コード": "B97", "店舗名": "練馬春日町店"},
    {"店舗コード": "B98", "店舗名": "船橋北本町店"},
    {"店舗コード": "003", "店舗名": "利根店"},
    {"店舗コード": "004", "店舗名": "北葛西店"},
    {"店舗コード": "007", "店舗名": "大森店"},
    {"店舗コード": "008", "店舗名": "柏松葉町店"},
    {"店舗コード": "009", "店舗名": "西国分寺店"},
    {"店舗コード": "011", "店舗名": "結城店"},
    {"店舗コード": "013", "店舗名": "千葉山王店"},
    {"店舗コード": "014", "店舗名": "小山店"},
    {"店舗コード": "016", "店舗名": "ユーカリが丘店"},
    {"店舗コード": "017", "店舗名": "福生熊川店"},
    {"店舗コード": "020", "店舗名": "野辺店"},
    {"店舗コード": "022", "店舗名": "久喜店"},
    {"店舗コード": "023", "店舗名": "瑞江店"},
    {"店舗コード": "025", "店舗名": "練馬氷川台店"},
    {"店舗コード": "026", "店舗名": "八千代店"},
    {"店舗コード": "028", "店舗名": "鳩ヶ谷里店"},
    {"店舗コード": "029", "店舗名": "稲毛園生店"},
    {"店舗コード": "030", "店舗名": "東大和店"},
    {"店舗コード": "031", "店舗名": "東川口店"},
    {"店舗コード": "032", "店舗名": "足立保木間店"},
    {"店舗コード": "033", "店舗名": "春日部店"},
    {"店舗コード": "034", "店舗名": "浦和西堀店"},
    {"店舗コード": "035", "店舗名": "川越今福店"},
    {"店舗コード": "036", "店舗名": "飯能緑町店"},
    {"店舗コード": "037", "店舗名": "川口芝店"},
    {"店舗コード": "038", "店舗名": "東鎌ヶ谷店"},
    {"店舗コード": "039", "店舗名": "境町店"},
    {"店舗コード": "040", "店舗名": "足立辰沼店"},
    {"店舗コード": "041", "店舗名": "船橋藤原店"},
    {"店舗コード": "042", "店舗名": "八王子上柚木店"},
    {"店舗コード": "043", "店舗名": "三郷店"},
    {"店舗コード": "044", "店舗名": "練馬西大泉店"},
    {"店舗コード": "045", "店舗名": "武蔵村山店"},
    {"店舗コード": "046", "店舗名": "狭山入曽店"},
    {"店舗コード": "047", "店舗名": "戸田本町店"},
    {"店舗コード": "049", "店舗名": "新松戸店"},
    {"店舗コード": "050", "店舗名": "練馬高松店"},
    {"店舗コード": "051", "店舗名": "川越旭町店"},
    {"店舗コード": "053", "店舗名": "青梅今寺店"},
    {"店舗コード": "054", "店舗名": "三郷戸ヶ崎店"},
    {"店舗コード": "055", "店舗名": "練馬中村橋店"},
    {"店舗コード": "056", "店舗名": "岩槻西町店"},
    {"店舗コード": "057", "店舗名": "千葉都町店"},
    {"店舗コード": "058", "店舗名": "松戸河原塚店"},
    {"店舗コード": "059", "店舗名": "八王子楢原店"},
    {"店舗コード": "060", "店舗名": "松戸五香店"},
    {"店舗コード": "061", "店舗名": "浦和三室店"},
    {"店舗コード": "062", "店舗名": "千葉みつわ台店"},
    {"店舗コード": "063", "店舗名": "千葉大宮台店"},
    {"店舗コード": "065", "店舗名": "船橋金杉店"},
    {"店舗コード": "066", "店舗名": "柏豊四季店"},
    {"店舗コード": "067", "店舗名": "松戸古ヶ崎店"},
    {"店舗コード": "069", "店舗名": "つくば竹園店"},
    {"店舗コード": "070", "店舗名": "松戸五香西店"},
    {"店舗コード": "071", "店舗名": "府中若松店"},
    {"店舗コード": "072", "店舗名": "足立鹿浜店"},
    {"店舗コード": "073", "店舗名": "下総中山店"},
    {"店舗コード": "101", "店舗名": "鎌ヶ谷店"},
    {"店舗コード": "102", "店舗名": "船橋山野町店"},
    {"店舗コード": "103", "店舗名": "足立竹の塚店"},
    {"店舗コード": "106", "店舗名": "行田店"},
    {"店舗コード": "107", "店舗名": "市原辰巳台店"},
    {"店舗コード": "108", "店舗名": "武蔵村山学園店"},
    {"店舗コード": "109", "店舗名": "成田店"},
    {"店舗コード": "110", "店舗名": "鶴ヶ島店"},
    {"店舗コード": "112", "店舗名": "水戸店"},
    {"店舗コード": "113", "店舗名": "八王子宇津木台店"},
    {"店舗コード": "114", "店舗名": "岩瀬店"},
    {"店舗コード": "117", "店舗名": "行方店"},
    {"店舗コード": "118", "店舗名": "富里店"},
    {"店舗コード": "120", "店舗名": "上尾店"},
    {"店舗コード": "121", "店舗名": "成田三里塚店"},
    {"店舗コード": "123", "店舗名": "白岡店"},
    {"店舗コード": "124", "店舗名": "取手東店"},
    {"店舗コード": "125", "店舗名": "吉川駅前通り店"},
    {"店舗コード": "127", "店舗名": "明野店"},
    {"店舗コード": "128", "店舗名": "蕨南町店"},
    {"店舗コード": "129", "店舗名": "つくば桜店"},
    {"店舗コード": "131", "店舗名": "水戸河和田店"},
    {"店舗コード": "132", "店舗名": "本庄店"},
    {"店舗コード": "133", "店舗名": "草加店"},
    {"店舗コード": "134", "店舗名": "壬生店"},
    {"店舗コード": "135", "店舗名": "前橋インターアカマル店"},
    {"店舗コード": "136", "店舗名": "柏西原店"},
    {"店舗コード": "137", "店舗名": "香取佐原店"},
    {"店舗コード": "138", "店舗名": "岩槻府内店"},
    {"店舗コード": "139", "店舗名": "笠間店"},
    {"店舗コード": "140", "店舗名": "東松山店"},
    {"店舗コード": "141", "店舗名": "東金店"},
    {"店舗コード": "142", "店舗名": "太田西本町店"},
    {"店舗コード": "143", "店舗名": "みどり笠懸店"},
    {"店舗コード": "144", "店舗名": "前橋駒形店"},
    {"店舗コード": "145", "店舗名": "鹿沼上殿店"},
    {"店舗コード": "146", "店舗名": "入間下藤沢店"},
    {"店舗コード": "147", "店舗名": "稲敷江戸崎店"},
    {"店舗コード": "148", "店舗名": "神栖波崎店"},
    {"店舗コード": "149", "店舗名": "熊谷石原店"},
    {"店舗コード": "150", "店舗名": "つくばみどりの店"},
    {"店舗コード": "151", "店舗名": "東村山青葉町店"},
    {"店舗コード": "152", "店舗名": "常陸太田宮本町店"},
    {"店舗コード": "153", "店舗名": "那須烏山店"},
    {"店舗コード": "154", "店舗名": "前橋北代田店"},
    {"店舗コード": "155", "店舗名": "沼田鍛冶町店"},
    {"店舗コード": "156", "店舗名": "香取多古町店"}
]
master_df = pd.DataFrame(STORE_MASTER_DATA)

# --- ログイン認証処理 ---
def check_password():
    if "password_correct" not in st.session_state:
        st.session_state["password_correct"] = False

    if st.session_state["password_correct"]:
        return True

    st.title("🔒 社内関係者専用ログイン")
    password_input = st.text_input("パスワードを入力してください", type="password")
    
    if st.button("ログイン"):
        if password_input == st.secrets.get("PASSWORD"):
            st.session_state["password_correct"] = True
            st.rerun()
        else:
            st.error("パスワードが正しくありません。")
    
    return False

if not check_password():
    st.stop()

# --- 店舗名自動補正関数 ---
def normalize_and_fix_store_names(df_results, master_df):
    if df_results.empty or master_df is None or "店舗コード" not in master_df.columns or "店舗名" not in master_df.columns:
        return df_results
    
    code_to_name = {}
    master_store_names = []
    
    for _, row in master_df.iterrows():
        raw_code = str(row["店舗コード"]).strip()
        raw_name = str(row["店舗名"]).strip().replace("ジェーソン", "")
        master_store_names.append(raw_name)
        
        code_to_name[raw_code] = raw_name
        if raw_code.isdigit():
            code_to_name[raw_code.zfill(3)] = raw_name
            code_to_name[str(int(raw_code))] = raw_name

    fixed_stores = []
    
    for idx, row in df_results.iterrows():
        fname = str(row.get("ファイル名", ""))
        ocr_store = str(row.get("店舗名", "")).strip()
        
        base_name = os.path.splitext(os.path.basename(fname))[0].strip()
        code_prefix = base_name.split('-')[0].split('_')[0].split(' ')[0].strip()
        
        matched_name = None
        
        if code_prefix in code_to_name:
            matched_name = code_to_name[code_prefix]
            
        if not matched_name:
            for code, name in code_to_name.items():
                if len(code) >= 2 and (base_name.startswith(code) or f"-{code}" in base_name or f"_{code}" in base_name):
                    matched_name = name
                    break
                    
        if not matched_name:
            clean_ocr = ocr_store.replace("ジェーソン", "").strip()
            if clean_ocr in master_store_names:
                matched_name = clean_ocr
        
        if matched_name:
            fixed_stores.append(f"ジェーソン{matched_name}")
        else:
            fixed_stores.append(ocr_store)
            
    df_results["店舗名"] = fixed_stores
    return df_results

# --- メイン画面 ---
st.title("📊 売上日計表 自動照合システム")
st.markdown("店舗からの売上日計表画像を自動照合し、全店舗マスタ（116店舗固定）と比較して**未提出店舗の抽出**およびExcel報告書を出力します。")

# SecretsからAPIキーを取得
api_key = st.secrets.get("GEMINI_API_KEY")

# 1. 調査対象日の設定（初期値：前日 / 上限：前日＝本日以降選択不可）
yesterday = datetime.date.today() - datetime.timedelta(days=1)

st.subheader("1. 調査対象日の設定")
col1, col2 = st.columns([1, 2])
with col1:
    ignore_date_check = st.checkbox("調査対象日を指定しない（すべての画像を取り込む）", value=False)
with col2:
    target_date = st.date_input(
        "調査対象日を選択してください", 
        value=yesterday, 
        max_value=yesterday,
        disabled=ignore_date_check
    )

target_date_str = "指定なし（全日付を照合対象とする）" if ignore_date_check else target_date.strftime("%Y/%m/%d")

# 2. 画像アップロード
st.subheader("2. 日計表画像のアップロード")
uploaded_files = st.file_uploader(
    "売上日計表の画像をアップロードしてください（複数選択可・順次追加可能）", 
    type=['png', 'jpg', 'jpeg'], 
    accept_multiple_files=True,
    key=f"uploader_{st.session_state['file_uploader_key']}"
)

# ボタンレイアウト
btn_col1, btn_col2, btn_col3 = st.columns([1.5, 1.5, 1.5])

with btn_col1:
    start_btn = st.button("🚀 照合を追加実行する", type="primary", use_container_width=True)

with btn_col2:
    clear_file_btn = st.button("📁 アップロード選択欄をクリア（結果は残す）", use_container_width=True)

with btn_col3:
    clear_all_btn = st.button("🗑️ 全調査完了（全結果データをクリアする）", use_container_width=True)

# アップロード枠のみクリア
if clear_file_btn:
    st.session_state["file_uploader_key"] += 1
    st.rerun()

# 全データクリア
if clear_all_btn:
    st.session_state["accumulated_results"] = pd.DataFrame()
    st.session_state["accumulated_unreadable"] = pd.DataFrame()
    st.session_state["file_uploader_key"] += 1
    st.success("🧹 照合結果データをすべてリセットしました。")
    st.rerun()

# 照合実行
if start_btn:
    if not api_key:
        st.error("システムエラー: APIキーが設定されていません。StreamlitのSecretsを設定してください。")
    elif not uploaded_files:
        st.warning("追加する画像を1枚以上アップロードしてください。")
    else:
        with st.spinner("画像を解析中です...（枚数によって数十秒〜数分かかります）"):
            try:
                genai.configure(api_key=api_key)
                model = genai.GenerativeModel(
                    'gemini-3.6-flash',
                    generation_config={"response_mime_type": "application/json"}
                )

                image_data_list = []
                file_name_mapping = []
                
                for idx, file in enumerate(uploaded_files):
                    img = PIL.Image.open(file)
                    image_data_list.append(img)
                    file_name_mapping.append(f"画像{idx+1}: {file.name}")
                
                mapping_text = "\n".join(file_name_mapping)

                master_info_str = "【店舗マスタ（ファイル名の数字/記号は店舗コードに対応しています）】\n"
                for _, r in master_df.iterrows():
                    master_info_str += f"- コード `{r['店舗コード']}`: {r['店舗名']}\n"

                prompt = f"""
                以下の売上日計表の画像（{len(uploaded_files)}枚）を読み取り、各画像について数値を照合し、JSONオブジェクト形式で出力してください。
                
                【アップロードされたファイル名と画像の順序】
                {mapping_text}
                
                {master_info_str}
                
                【重要：店舗名の判定規則】
                ・ファイル名の先頭部分（例: `B75.jpg` -> `B75`, `144-1.jpg` -> `144`）は店舗マスタの店舗コードに対応しています。
                ・店舗名は必ず店舗マスタに記載された正確な店舗名（例: `ジェーソン和光店`）を判定してください。
                
                【重要：判別不能・不明画像の無視（スキップ）ルール】
                ・画像がブレている、極度に不鮮明、見切れている、または売上日計表ではないなどの理由で数値や店舗名が判別できない画像は照合対象から除外してください。
                ・除外した画像のファイル名は `unreadable_files` 配列に記載してください。
                
                【重要：調査対象日フィルター規則】
                ・指定の調査対象日：{target_date_str}
                ・「{target_date_str}」が「指定なし」でない場合、画像内の売上日計表の日付（年月日）が対象日と一致するか判定してください。
                ・日付が一致しない場合、各照合項目を「対象外(日付不一致)」とし、不一致の内容に「指定された調査対象日（{target_date_str}）と異なります」と記録してください。
                
                【基本照合項目（左側の「【売上集計表】レジ精算総計」と、右側の集計欄の数値を比較）】
                1.現金在高  2.過不足  3.商品券  4.キャッシュレス  5.入金額
                
                【追加の照合条件（ミスパンチ・返金）】
                ・売上日計表に記載されている「ミスパンチ」および「返金」の金額を確認してください。
                ・「ミスパンチ」の金額が0円でない場合：「返金ミスパンチ貼付用紙」の「ミスパンチ合計」と数値を照合してください。
                ・「返金」の金額が0円でない場合：「返金ミスパンチ貼付用紙」の「返金合計」と数値を照合してください。
                ・金額が0円または記載がない場合は、ステータスを「対象外(0円)」としてください。
                
                【出力フォーマット】
                {{
                  "results": [
                    {{
                      "ファイル名": "...", 
                      "店舗名": "店舗名", 
                      "日付": "YYYY/MM/DD", 
                      "現金在高": "一致/不一致/対象外(日付不一致)",
                      "過不足": "一致/不一致/対象外(日付不一致)",
                      "商品券": "一致/不一致/対象外(日付不一致)",
                      "キャッシュレス": "一致/不一致/対象外(日付不一致)",
                      "入金額": "一致/不一致/対象外(日付不一致)",
                      "ミスパンチ照合": "一致/不一致/対象外(0円)/対象外(日付不一致)",
                      "返金照合": "一致/不一致/対象外(0円)/対象外(日付不一致)",
                      "不一致の内容": "詳細（すべて一致・対象外の場合は 'なし'）"
                    }}
                  ],
                  "unreadable_files": ["判別不能なファイル名1.jpg", "..."]
                }}
                """

                contents = [prompt] + image_data_list
                response = model.generate_content(contents)
                raw_json = json.loads(response.text)

                if isinstance(raw_json, dict):
                    result_data = raw_json.get("results", [])
                    unreadable_files = raw_json.get("unreadable_files", [])
                else:
                    result_data = raw_json
                    unreadable_files = []

                new_df_results = pd.DataFrame(result_data)
                
                if not new_df_results.empty:
                    new_df_results = normalize_and_fix_store_names(new_df_results, master_df)
                    for col in new_df_results.columns:
                        new_df_results[col] = new_df_results[col].astype(str)

                new_df_unreadable = pd.DataFrame({"判別不能ファイル名": unreadable_files, "理由": "画像不鮮明・判別不能につき無視"}) if unreadable_files else pd.DataFrame()
                if not new_df_unreadable.empty:
                    for col in new_df_unreadable.columns:
                        new_df_unreadable[col] = new_df_unreadable[col].astype(str)

                # 追記と重複排除
                if not new_df_results.empty:
                    combined_res = pd.concat([st.session_state["accumulated_results"], new_df_results], ignore_index=True)
                    st.session_state["accumulated_results"] = combined_res.drop_duplicates(subset=["ファイル名"], keep="first")

                if not new_df_unreadable.empty:
                    combined_unread = pd.concat([st.session_state["accumulated_unreadable"], new_df_unreadable], ignore_index=True)
                    st.session_state["accumulated_unreadable"] = combined_unread.drop_duplicates(subset=["判別不能ファイル名"], keep="first")

                st.success("✅ 解析結果を追加・更新しました！")

            except Exception as e:
                st.error(f"システムエラーが発生しました: {e}")

# 画面表示部
df_results = st.session_state["accumulated_results"]
df_unreadable = st.session_state["accumulated_unreadable"]

if not df_results.empty or not df_unreadable.empty:
    df_unsubmitted = pd.DataFrame()
    if not df_results.empty and ("現金を高" in df_results.columns or "現金在高" in df_results.columns):
        col_name = "現金在高" if "現金在高" in df_results.columns else "現金を高"
        valid_results = df_results[~df_results[col_name].astype(str).str.contains("日付不一致")]
        extracted_stores = valid_results["店舗名"].astype(str).str.replace("ジェーソン", "").str.strip().unique()
    else:
        extracted_stores = []

    df_unsubmitted = master_df[~master_df["店舗名"].astype(str).str.strip().isin(extracted_stores)].copy()
    df_unsubmitted["提出ステータス"] = "未提出"
    for col in df_unsubmitted.columns:
        df_unsubmitted[col] = df_unsubmitted[col].astype(str)

    tab1, tab2, tab3 = st.tabs([
        f"📊 照合結果（累計: {len(df_results)}件）", 
        f"⚠️ 未提出店舗一覧（{len(df_unsubmitted)}店舗）" if not df_unsubmitted.empty else "⚠️ 未提出店舗一覧",
        f"🚫 判別不能・無視ファイル（{len(df_unreadable)}件）" if not df_unreadable.empty else "🚫 判別不能・無視ファイル"
    ])

    def highlight_mismatch(val):
        if val == '不一致':
            return 'background-color: #ffcccc'
        elif '日付不一致' in str(val):
            return 'background-color: #eeeeee; color: #888888'
        return ''

    with tab1:
        st.subheader(f"提出分・累計数値照合結果（対象日: {target_date_str}）")
        if not df_results.empty:
            st.dataframe(df_results.style.map(highlight_mismatch), use_container_width=True)

    with tab2:
        st.subheader("未提出店舗リスト")
        if not df_unsubmitted.empty:
            st.warning(f"全{len(master_df)}店舗中、対象日（{target_date_str}）の画像が **{len(df_unsubmitted)}店舗** 未提出です。")
            st.dataframe(df_unsubmitted, use_container_width=True)
        else:
            st.info("すべての対象店舗の画像が提出されています。")

    with tab3:
        st.subheader("判別不能・無視されたファイル")
        if not df_unreadable.empty:
            st.error(f"以下の **{len(df_unreadable)}件** の画像は画質不良または不整合のため判別できず除外されています。")
            st.dataframe(df_unreadable, use_container_width=True)
        else:
            st.success("判別不能・無視された画像はありません。")

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        if not df_results.empty:
            df_results.to_excel(writer, index=False, sheet_name="照合結果")
        if not df_unsubmitted.empty:
            df_unsubmitted.to_excel(writer, index=False, sheet_name="未提出店舗一覧")
        if not df_unreadable.empty:
            df_unreadable.to_excel(writer, index=False, sheet_name="判別不能ファイル一覧")

    output.seek(0)
    wb = openpyxl.load_workbook(output)

    header_fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alignment_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), 
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    gray_fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
    alt_fill = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = alignment_center
            cell.border = thin_border

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=2):
            for col_idx, cell in enumerate(row):
                cell.border = thin_border
                if sheet_name == "照合結果" and col_idx not in [0, 1, 10]:
                    cell.alignment = alignment_center
                else:
                    cell.alignment = alignment_left

                if row_idx % 2 == 0:
                    cell.fill = alt_fill

                if cell.value == "不一致":
                    cell.fill = red_fill
                elif cell.value == "対象外(日付不一致)":
                    cell.fill = gray_fill

    final_output = io.BytesIO()
    wb.save(final_output)
    excel_data = final_output.getvalue()

    st.download_button(
        label="📥 累計照合結果・未提出・不明ファイルリスト（Excel）をダウンロード",
        data=excel_data,
        file_name="売上照合および未提出店舗報告書_累計.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
